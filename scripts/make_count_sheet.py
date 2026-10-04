import json, re
from openpyxl import Workbook
from datetime import date
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]  # the WMS folder

D = json.load(open(ROOT / 'data' / 'generated' / 'count.json'))
OUT = str(ROOT / "outputs" / "count-sheets" / f"WAREHOUSE_STOCK_COUNT_{date.today():%Y%m%d}.xlsx")
MAXROW = 1000          # dropdowns cover rows 2..1000 so inserted rows (mixed pallets) keep them
FLOOR_ROWS = 10        # blank floor rows per rack side

products = [p for p in D['products'] if p['b'] != 'Deccan' and 'USE FOR NEW PRODUCT' not in p['k'].upper()]

thin = Side(style='thin'); border = Border(left=thin, right=thin, top=thin, bottom=thin)
hfont = Font(name='Calibri', size=11, bold=True); font = Font(name='Calibri', size=11)
hfill = PatternFill('solid', fgColor='FFB7B7B7')
center = Alignment(horizontal='center', vertical='center')
HDR = ["Location", "Item", "No. of Cases", "Best Before Date", "Batch No", "Container"]
W = [11.6, 53, 6.4, 12.3, 19, 21.7]

wb = Workbook(); wb.remove(wb.active)

# ---------- PRODUCT LIST (dropdown source + counted totals)
pl = wb.create_sheet("PRODUCT LIST")
pl_hdr = ["#", "CATEGORY", "BRAND", "PRODUCT CODE", "PRODUCT", "PRODUCT DESC", "PKG SIZE", "CASES / PALLET", "COUNTED CASES"]
for i, (h, w) in enumerate(zip(pl_hdr, [5, 14, 16, 13, 60, 55, 16, 10, 12]), 1):
    c = pl.cell(1, i, h); c.font = hfont; c.fill = hfill; c.border = border; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    pl.column_dimensions[L(i)].width = w
pl.row_dimensions[1].height = 30
pl.freeze_panes = "A2"
NP = len(products)
LIST_REF = f"'PRODUCT LIST'!$E$2:$E${NP + 1}"

def style_row(ws, r, c0, n, fmt=True):
    for i in range(n):
        cell = ws.cell(r, c0 + i); cell.border = border; cell.font = font
    ws.cell(r, c0 + 2).alignment = Alignment(horizontal='center')
    if fmt:
        ws.cell(r, c0 + 3).number_format = 'mmm yyyy'; ws.cell(r, c0 + 3).alignment = Alignment(horizontal='center')
        ws.cell(r, c0 + 5).number_format = '0'

count_ranges = []   # (sheet, item col letter, cases col letter) for the totals formula

def block_sheet(name, blocks, header=HDR, fmt=True):
    ws = wb.create_sheet(name)
    for b in range(blocks):
        c0 = 1 + b * 7
        for i, h in enumerate(header):
            c = ws.cell(1, c0 + i, h); c.font = hfont; c.fill = hfill; c.border = border; c.alignment = center
            ws.column_dimensions[L(c0 + i)].width = (W if b == 0 else [11.6, 53, 6.4, 12.3, 19.3, 14.3])[i] if i < 6 else 30
        if b < blocks - 1: ws.column_dimensions[L(c0 + 6)].width = 5
        item, cases = L(c0 + 1), L(c0 + 2)
        dv = DataValidation(type="list", formula1="=" + LIST_REF, allow_blank=True, showDropDown=False,
                            showErrorMessage=True, errorStyle="warning", errorTitle="Not in the product list",
                            error="This product isn't in the PRODUCT LIST tab. Keep it anyway?")
        dv.add(f"{item}2:{item}{MAXROW}"); ws.add_data_validation(dv)
        nv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True,
                            showErrorMessage=True, errorStyle="stop", errorTitle="Cases", error="Enter a whole number of cases.")
        nv.add(f"{cases}2:{cases}{MAXROW}"); ws.add_data_validation(nv)
        count_ranges.append((name, item, cases))
    ws.freeze_panes = "A2"
    return ws

def write_lists(ws, start, lists, fmt=True):
    r = start
    for i in range(max([0] + [len(l) for l in lists])):
        for s, l in enumerate(lists):
            if i < len(l):
                c0 = 1 + s * 7
                ws.cell(r, c0, l[i]); style_row(ws, r, c0, 6, fmt)
        r += 1
    return r

pad = lambda n: f"{n:02d}"
faces = D['faces']
def rack_of(fid): return ('D' + fid[1:-1]) if fid.startswith('D') else fid[:-1]
def slot_list(face, c):
    col = next((x for x in face['cols'] if x[0] == c), None)
    if not col: return []
    return [f"{face['id']}-{pad(c)}-L{l}-B{b}" for l, b in sorted(col[1], key=lambda s: (s[1], s[0]))]

def rack_sheet(rack):
    sides = sorted([f for f in faces if rack_of(f['id']) == rack], key=lambda f: f['id'][-1])
    name = "Rack00" if rack == "0" else ("DOOR" + rack[1:] if rack.startswith("D") else "RACK" + pad(int(rack)))
    ws = block_sheet(name, len(sides))
    cols = sorted({c for f in sides for c, _ in f['cols']})
    r = 2
    for c in cols:
        r = write_lists(ws, r, [slot_list(f, c) for f in sides]) + 1
    r = write_lists(ws, r, [[f"{f['id']} Floor"] * FLOOR_ROWS for f in sides])
    if rack == "7": write_lists(ws, r, [["Lane 7"] * 5])

def area_sheet(name, header, blocks, rows, fmt=True):
    ws = block_sheet(name, blocks, header, fmt)
    write_lists(ws, 2, [rows], fmt)

# Same tab order as the stock workbooks
area_sheet("Anaganaga", HDR + ["Description"], 1, ["Between line 2 & 3"] * 40)
wb["Anaganaga"].column_dimensions['G'].width = 30
area_sheet("NearCooler", HDR, 1, [f"NEARCOOLER-0{n}" for n in range(1, 6) for _ in range(10)])
racks = list(dict.fromkeys(rack_of(f['id']) for f in faces))
for r in [x for x in racks if x.startswith('D')]: rack_sheet(r)
area_sheet("DOOR2N3", HDR, 2, ["DOOR 2N3"] * 10)
area_sheet("Front", HDR, 1, ["LOAD_AREA_01"] * 20)
area_sheet("CSFloor", ["Location", "Item", "QTY", "Container", "WH Arrival Date", "Batch No"], 2, ["Recv Floor"] * 20, fmt=False)
for r in [x for x in racks if not x.startswith('D')]: rack_sheet(r)

# Product rows + a live total of counted cases across every tab
for i, p in enumerate(products, 2):
    pl.cell(i, 1, i - 1); pl.cell(i, 2, p['cat'].upper()); pl.cell(i, 3, p['b'].upper()); pl.cell(i, 4, p['c'])
    pl.cell(i, 5, p['k']); pl.cell(i, 6, p['n']); pl.cell(i, 7, p['pk']); pl.cell(i, 8, p['pl'])
    terms = [f"SUMIF('{s}'!${it}$2:${it}${MAXROW},$E{i},'{s}'!${ca}$2:${ca}${MAXROW})" for s, it, ca in count_ranges]
    pl.cell(i, 9, "=" + "+".join(terms))
    for c in range(1, 10): pl.cell(i, c).border = border; pl.cell(i, c).font = font
    for c in (1, 4, 8, 9): pl.cell(i, c).alignment = Alignment(horizontal='center')
pl.auto_filter.ref = f"A1:I{NP + 1}"
print('formula length', len(pl.cell(2, 9).value), 'tabs', len(wb.sheetnames), 'products', NP)
wb.save(OUT)
print(wb.sheetnames)
